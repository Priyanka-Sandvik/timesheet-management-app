"""Unit tests: xlsx workbook structure built by export_service, re-read via openpyxl."""
from __future__ import annotations

import io
from datetime import date, datetime, timezone

import openpyxl
import pytest

from app.models.entities import TimesheetEntity
from app.services.export_service import (
    EmployeeMonthData,
    build_monthly_workbook,
    sanitize_sheet_name,
)


def _entity(user_id: str, entry_date: date, task_code: str, hours: float) -> TimesheetEntity:
    now = datetime.now(timezone.utc)
    return TimesheetEntity(
        user_id=user_id,
        entry_date=entry_date,
        task_code=task_code,
        task_name=f"{task_code} name",
        project_name="Project Z",
        week_start_date=entry_date,
        hours_logged=hours,
        notes=None,
        status="Submitted",
        created_at=now,
        updated_at=now,
    )


def test_sanitize_sheet_name_strips_invalid_characters_and_truncates() -> None:
    used: set[str] = set()
    name = sanitize_sheet_name("a/very*long[email]:address?that\\exceeds_31_chars@sandvik.com", used)
    assert len(name) <= 31
    for ch in "\\/*[]:?":
        assert ch not in name


def test_sanitize_sheet_name_deduplicates_collisions() -> None:
    used: set[str] = set()
    n1 = sanitize_sheet_name("alice@sandvik.com", used)
    n2 = sanitize_sheet_name("alice@sandvik.com", used)
    assert n1 != n2
    assert len(n2) <= 31


def test_workbook_has_one_sheet_per_employee_including_zero_hour() -> None:
    employees = [
        EmployeeMonthData(
            email="alice@sandvik.com",
            entries=[_entity("alice@sandvik.com", date(2026, 8, 3), "TASK-A", 8.0)],
        ),
        EmployeeMonthData(email="bob@sandvik.com", entries=[]),  # zero-hour employee
    ]
    wb_bytes = build_monthly_workbook("2026-08", employees)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))

    assert set(wb.sheetnames) == {"Summary", "alice@sandvik.com", "bob@sandvik.com"}


def test_zero_hour_employee_sheet_has_no_entries_placeholder() -> None:
    employees = [EmployeeMonthData(email="bob@sandvik.com", entries=[])]
    wb_bytes = build_monthly_workbook("2026-08", employees)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))

    ws = wb["bob@sandvik.com"]
    assert ws.cell(row=2, column=1).value == "No entries"


def test_sheet_layout_columns_and_totals() -> None:
    entries = [
        _entity("alice@sandvik.com", date(2026, 8, 3), "TASK-A", 4.0),
        _entity("alice@sandvik.com", date(2026, 8, 4), "TASK-A", 4.0),
        _entity("alice@sandvik.com", date(2026, 8, 3), "TASK-B", 2.0),
    ]
    employees = [EmployeeMonthData(email="alice@sandvik.com", entries=entries)]
    wb_bytes = build_monthly_workbook("2026-08", employees)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb["alice@sandvik.com"]

    # Header row: Task | 1 | 2 | ... | 31 | Total
    assert ws.cell(row=1, column=1).value == "Task"
    assert ws.cell(row=1, column=2).value == 1
    assert ws.cell(row=1, column=32).value == 31
    assert ws.cell(row=1, column=33).value == "Total"

    # Task A row: day 3 -> 4.0, day 4 -> 4.0, total 8.0
    task_a_row = 2
    assert ws.cell(row=task_a_row, column=1).value == "TASK-A name"
    assert ws.cell(row=task_a_row, column=1 + 3).value == 4.0  # day-of-month 3 -> column 1+3
    assert ws.cell(row=task_a_row, column=1 + 4).value == 4.0
    assert ws.cell(row=task_a_row, column=33).value == 8.0

    # Task B row: day 3 -> 2.0, total 2.0
    task_b_row = 3
    assert ws.cell(row=task_b_row, column=1).value == "TASK-B name"
    assert ws.cell(row=task_b_row, column=1 + 3).value == 2.0
    assert ws.cell(row=task_b_row, column=33).value == 2.0

    # Grand total row should exist somewhere below the task+weekly-subtotal rows.
    grand_total_row = None
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "Grand Total":
            grand_total_row = row
            break
    assert grand_total_row is not None
    # Aug 3 total across tasks = 4.0 + 2.0 = 6.0; Aug 4 total = 4.0.
    assert ws.cell(row=grand_total_row, column=1 + 3).value == 6.0
    assert ws.cell(row=grand_total_row, column=1 + 4).value == 4.0
    assert ws.cell(row=grand_total_row, column=33).value == 10.0


def test_summary_sheet_is_first_and_lists_every_active_employee() -> None:
    employees = [
        EmployeeMonthData(
            email="alice@sandvik.com",
            entries=[_entity("alice@sandvik.com", date(2026, 8, 3), "TASK-A", 8.0)],
        ),
        EmployeeMonthData(email="bob@sandvik.com", entries=[]),
    ]
    wb_bytes = build_monthly_workbook("2026-08", employees)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))

    assert wb.sheetnames[0] == "Summary"
    ws = wb["Summary"]
    assert ws.cell(row=1, column=1).value == "Employee"
    assert ws.cell(row=1, column=2).value == 1
    assert ws.cell(row=1, column=32).value == 31
    assert ws.cell(row=1, column=33).value == "Total"
    # Rows sorted by email, same order as the per-employee sheets.
    assert ws.cell(row=2, column=1).value == "alice@sandvik.com"
    assert ws.cell(row=3, column=1).value == "bob@sandvik.com"


def test_summary_sheet_per_employee_daily_and_monthly_totals() -> None:
    employees = [
        EmployeeMonthData(
            email="alice@sandvik.com",
            entries=[
                _entity("alice@sandvik.com", date(2026, 8, 3), "TASK-A", 4.0),
                _entity("alice@sandvik.com", date(2026, 8, 3), "TASK-B", 2.0),
                _entity("alice@sandvik.com", date(2026, 8, 4), "TASK-A", 4.0),
            ],
        ),
        EmployeeMonthData(
            email="bob@sandvik.com",
            entries=[_entity("bob@sandvik.com", date(2026, 8, 3), "TASK-A", 1.5)],
        ),
    ]
    wb_bytes = build_monthly_workbook("2026-08", employees)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb["Summary"]

    # Alice: Aug 3 = 4.0 + 2.0 = 6.0, Aug 4 = 4.0, monthly total = 10.0.
    assert ws.cell(row=2, column=1 + 3).value == 6.0
    assert ws.cell(row=2, column=1 + 4).value == 4.0
    assert ws.cell(row=2, column=33).value == 10.0

    # Bob: Aug 3 = 1.5, monthly total = 1.5.
    assert ws.cell(row=3, column=1 + 3).value == 1.5
    assert ws.cell(row=3, column=33).value == 1.5

    # Grand Total row: Aug 3 = 6.0 + 1.5 = 7.5, Aug 4 = 4.0, monthly total = 11.5.
    assert ws.cell(row=4, column=1).value == "Grand Total"
    assert ws.cell(row=4, column=1 + 3).value == 7.5
    assert ws.cell(row=4, column=1 + 4).value == 4.0
    assert ws.cell(row=4, column=33).value == 11.5
