"""CSV/XLSX parsing + validation for POST /admin/tasks/import.

Rules (architecture doc §3, resolved open question 1):
- Columns: TaskCode (required, unique within file), TaskName (required), Description
  (optional), Sponsor (optional), CostCentre (optional), CoEResponsible (optional),
  Status (Open/OnHold/Closed, default Open).
- Header row required. Unknown extra columns ignored.
- Missing required COLUMNS (TaskCode or TaskName header absent entirely) fails the WHOLE
  import - no rows processed - and raises an AppError describing the missing column(s).
- Per-row validation errors (missing TaskCode value, duplicate TaskCode within the file,
  invalid Status value) are captured per-row in `errors` and that row is skipped; OTHER
  valid rows still import.
- .csv parsed with the stdlib `csv` module; .xlsx parsed with `openpyxl` in read-only mode,
  loaded entirely from in-memory bytes - never written to disk.
"""
from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import load_workbook

from py_common.core.errors import AppError

from app.models.schemas import VALID_STATUSES, ImportParseResult, ImportRowError, ParsedTaskRow

REQUIRED_COLUMNS = ("TaskCode", "TaskName")
KNOWN_COLUMNS = ("TaskCode", "TaskName", "Description", "Sponsor", "CostCentre", "CoEResponsible", "Status")

# Alternate header spellings accepted for known columns (case-insensitive, matched after
# stripping whitespace). Keeps the import file human-friendly (e.g. "Cost Centre" with a
# space) while KNOWN_COLUMNS stays the canonical internal name.
HEADER_ALIASES: dict[str, str] = {
    "task name": "TaskName",
    "taskname": "TaskName",
    "task code": "TaskCode",
    "taskcode": "TaskCode",
    "description": "Description",
    "sponsor": "Sponsor",
    "cost centre": "CostCentre",
    "cost center": "CostCentre",
    "costcentre": "CostCentre",
    "costcenter": "CostCentre",
    "coe responsible": "CoEResponsible",
    "coeresponsible": "CoEResponsible",
    "coe": "CoEResponsible",
    "status": "Status",
}

DEFAULT_STATUS = "Open"


def _normalize_header(raw_headers: list[Any]) -> dict[str, int]:
    """Maps a known column name -> its column index, matched case-insensitively and
    trimmed of surrounding whitespace (including common alternate spellings). Unknown
    columns are ignored (not included).
    """
    mapping: dict[str, int] = {}
    for idx, raw in enumerate(raw_headers):
        if raw is None:
            continue
        name = str(raw).strip()
        known = HEADER_ALIASES.get(name.lower())
        if known and known not in mapping:
            mapping[known] = idx
    return mapping


def _check_required_columns(header_map: dict[str, int]) -> None:
    missing = [c for c in REQUIRED_COLUMNS if c not in header_map]
    if missing:
        raise AppError(
            status_code=400,
            code="VALIDATION_ERROR",
            message=f"Import file is missing required column(s): {', '.join(missing)}",
            details=[{"field": c, "issue": "required column missing from header row"} for c in missing],
        )


def _cell(row: list[Any], header_map: dict[str, int], column: str) -> str | None:
    idx = header_map.get(column)
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _validate_row(row_number: int, raw_row: list[Any], header_map: dict[str, int], seen_task_codes: set[str]):
    """Returns (ParsedTaskRow, None) on success or (None, ImportRowError) on failure."""
    task_code = _cell(raw_row, header_map, "TaskCode")
    if not task_code:
        return None, ImportRowError(row=row_number, reason="Missing required value for TaskCode")

    if task_code in seen_task_codes:
        return None, ImportRowError(row=row_number, reason=f"Duplicate TaskCode '{task_code}' within file")

    task_name = _cell(raw_row, header_map, "TaskName")
    if not task_name:
        return None, ImportRowError(row=row_number, reason="Missing required value for TaskName")

    description = _cell(raw_row, header_map, "Description")
    sponsor = _cell(raw_row, header_map, "Sponsor")
    cost_centre = _cell(raw_row, header_map, "CostCentre")
    coe_responsible = _cell(raw_row, header_map, "CoEResponsible")

    status_raw = _cell(raw_row, header_map, "Status")
    status = status_raw if status_raw is not None else DEFAULT_STATUS
    if status not in VALID_STATUSES:
        return None, ImportRowError(row=row_number, reason=f"Invalid Status value '{status}'")

    seen_task_codes.add(task_code)
    return (
        ParsedTaskRow(
            task_code=task_code,
            task_name=task_name,
            description=description,
            sponsor=sponsor,
            cost_centre=cost_centre,
            coe_responsible=coe_responsible,
            status=status,
        ),
        None,
    )


def parse_csv(content: bytes) -> ImportParseResult:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        raise AppError(
            status_code=400,
            code="VALIDATION_ERROR",
            message="Import file is empty; a header row is required",
        )

    header_map = _normalize_header(header)
    _check_required_columns(header_map)

    rows: list[ParsedTaskRow] = []
    errors: list[ImportRowError] = []
    seen: set[str] = set()

    for offset, raw_row in enumerate(reader):
        row_number = offset + 2  # +1 for 1-indexing, +1 for the header row
        if not any(str(c).strip() for c in raw_row if c is not None):
            continue  # skip fully blank lines silently
        parsed, error = _validate_row(row_number, raw_row, header_map, seen)
        if error is not None:
            errors.append(error)
        else:
            rows.append(parsed)

    return ImportParseResult(rows=rows, errors=errors)


def parse_xlsx(content: bytes) -> ImportParseResult:
    workbook = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = list(next(rows_iter))
        except StopIteration:
            raise AppError(
                status_code=400,
                code="VALIDATION_ERROR",
                message="Import file is empty; a header row is required",
            )

        header_map = _normalize_header(header)
        _check_required_columns(header_map)

        rows: list[ParsedTaskRow] = []
        errors: list[ImportRowError] = []
        seen: set[str] = set()

        for offset, raw_row in enumerate(rows_iter):
            row_number = offset + 2
            raw_row_list = list(raw_row)
            if not any(c is not None and str(c).strip() for c in raw_row_list):
                continue  # skip fully blank rows silently
            parsed, error = _validate_row(row_number, raw_row_list, header_map, seen)
            if error is not None:
                errors.append(error)
            else:
                rows.append(parsed)

        return ImportParseResult(rows=rows, errors=errors)
    finally:
        workbook.close()


def parse_import_file(filename: str, content: bytes) -> ImportParseResult:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return parse_csv(content)
    if lower.endswith(".xlsx"):
        return parse_xlsx(content)
    raise AppError(
        status_code=400,
        code="VALIDATION_ERROR",
        message="Unsupported file type; only .csv and .xlsx are accepted",
        details=[{"field": "file", "issue": f"unsupported extension for '{filename}'"}],
    )
